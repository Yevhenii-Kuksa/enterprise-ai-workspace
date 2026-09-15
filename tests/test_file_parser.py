from io import BytesIO

import pytest
from app.ingestion.file_parser import (
    DOCX_MIME_TYPE,
    FileParseRequest,
    parse_document_file,
    resolve_document_type,
    validate_mime_type,
)
from app.ingestion.parsing import SupportedDocumentType
from docx import Document as DocxDocument


def test_resolve_document_type_supports_known_extensions() -> None:
    assert (
        resolve_document_type("instrukcja.txt")
        is SupportedDocumentType.TXT
    )
    assert (
        resolve_document_type("procedura.PDF")
        is SupportedDocumentType.PDF
    )
    assert (
        resolve_document_type("regulamin.docx")
        is SupportedDocumentType.DOCX
    )


def test_resolve_document_type_rejects_empty_filename() -> None:
    with pytest.raises(
        ValueError,
        match="Filename must not be empty.",
    ):
        resolve_document_type("   ")


def test_resolve_document_type_rejects_missing_extension() -> None:
    with pytest.raises(
        ValueError,
        match="Unsupported document extension: <none>.",
    ):
        resolve_document_type("README")


def test_resolve_document_type_rejects_unsupported_extension() -> None:
    with pytest.raises(
        ValueError,
        match=r"Unsupported document extension: \.xlsx\.",
    ):
        resolve_document_type("raport.xlsx")


def test_validate_mime_type_uses_expected_type_when_missing() -> None:
    assert (
        validate_mime_type(
            SupportedDocumentType.PDF,
            None,
        )
        == "application/pdf"
    )


def test_validate_mime_type_rejects_empty_value() -> None:
    with pytest.raises(
        ValueError,
        match="MIME type must not be empty.",
    ):
        validate_mime_type(
            SupportedDocumentType.PDF,
            "   ",
        )


def test_validate_mime_type_rejects_extension_mismatch() -> None:
    with pytest.raises(
        ValueError,
        match="MIME type does not match document extension.",
    ):
        validate_mime_type(
            SupportedDocumentType.PDF,
            "text/plain",
        )


def test_parse_document_file_dispatches_txt() -> None:
    parsed = parse_document_file(
        FileParseRequest(
            filename="instrukcja.txt",
            content=b"Instrukcja stanowiskowa.",
            mime_type="text/plain",
        )
    )

    assert parsed.document_type is SupportedDocumentType.TXT
    assert parsed.text == "Instrukcja stanowiskowa."


def test_parse_document_file_dispatches_docx() -> None:
    buffer = BytesIO()

    document = DocxDocument()
    document.add_heading(
        "Procedura",
        level=1,
    )
    document.add_paragraph(
        "Treść procedury."
    )
    document.save(buffer)

    parsed = parse_document_file(
        FileParseRequest(
            filename="procedura.docx",
            content=buffer.getvalue(),
            mime_type=DOCX_MIME_TYPE,
        )
    )

    assert parsed.document_type is SupportedDocumentType.DOCX
    assert len(parsed.parts) == 1
    assert parsed.parts[0].text == "Treść procedury."
    assert parsed.parts[0].section_title == "Procedura"


def test_parse_document_file_rejects_empty_content() -> None:
    with pytest.raises(
        ValueError,
        match="Document content must not be empty.",
    ):
        parse_document_file(
            FileParseRequest(
                filename="instrukcja.txt",
                content=b"",
                mime_type="text/plain",
            )
        )


def test_parse_document_file_rejects_mime_mismatch_before_parsing() -> None:
    with pytest.raises(
        ValueError,
        match="MIME type does not match document extension.",
    ):
        parse_document_file(
            FileParseRequest(
                filename="fake.pdf",
                content=b"not-even-a-pdf",
                mime_type="text/plain",
            )
        )