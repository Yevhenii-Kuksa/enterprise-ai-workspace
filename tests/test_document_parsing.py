from io import BytesIO

import pytest
from app.ingestion.chunking import ChunkingConfig
from app.ingestion.parsing import (
    ParsedDocument,
    ParsedDocumentPart,
    SupportedDocumentType,
    chunk_parsed_document,
    parse_docx_document,
    parse_pdf_document,
    parse_txt_document,
)
from docx import Document as DocxDocument
from pypdf import PdfWriter


def test_parse_txt_document_returns_normalized_document() -> None:
    parsed = parse_txt_document(
        b"Pierwsza linia.\r\nDruga linia."
    )

    assert parsed.document_type is SupportedDocumentType.TXT
    assert parsed.mime_type == "text/plain"
    assert parsed.text == "Pierwsza linia.\nDruga linia."
    assert len(parsed.parts) == 1

    part = parsed.parts[0]

    assert part.text == "Pierwsza linia.\nDruga linia."
    assert part.page_number is None
    assert part.section_title is None
    assert part.source_locator == {
        "type": "document",
    }


def test_parse_txt_document_rejects_empty_content() -> None:
    with pytest.raises(
        ValueError,
        match="Document content must not be empty.",
    ):
        parse_txt_document(b"")


def test_parse_txt_document_rejects_whitespace_only_content() -> None:
    with pytest.raises(
        ValueError,
        match="TXT document does not contain readable text.",
    ):
        parse_txt_document(b"   \n\r\n   ")


def test_parse_txt_document_rejects_invalid_utf8() -> None:
    with pytest.raises(
        ValueError,
        match="TXT document is not valid utf-8 text.",
    ):
        parse_txt_document(b"\xff\xfe\xfa")


def test_parse_pdf_document_rejects_empty_content() -> None:
    with pytest.raises(
        ValueError,
        match="Document content must not be empty.",
    ):
        parse_pdf_document(b"")


def test_parse_pdf_document_rejects_invalid_pdf() -> None:
    with pytest.raises(
        ValueError,
        match="PDF document could not be parsed.",
    ):
        parse_pdf_document(b"not-a-pdf")


def test_parse_pdf_document_rejects_pdf_without_text() -> None:
    buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(
        width=100,
        height=100,
    )
    writer.write(buffer)

    with pytest.raises(
        ValueError,
        match="PDF document does not contain extractable text.",
    ):
        parse_pdf_document(buffer.getvalue())


def test_parse_docx_document_preserves_heading_metadata() -> None:
    buffer = BytesIO()

    document = DocxDocument()
    document.add_heading(
        "Procedura magazynowa",
        level=1,
    )
    document.add_paragraph(
        "Towar należy sprawdzić przed przyjęciem."
    )
    document.add_heading(
        "Kontrola jakości",
        level=2,
    )
    document.add_paragraph(
        "Należy zweryfikować stan opakowania."
    )
    document.save(buffer)

    parsed = parse_docx_document(
        buffer.getvalue()
    )

    assert parsed.document_type is SupportedDocumentType.DOCX
    assert len(parsed.parts) == 2

    assert parsed.parts[0].text == (
        "Towar należy sprawdzić przed przyjęciem."
    )
    assert (
        parsed.parts[0].section_title
        == "Procedura magazynowa"
    )

    assert parsed.parts[1].text == (
        "Należy zweryfikować stan opakowania."
    )
    assert (
        parsed.parts[1].section_title
        == "Kontrola jakości"
    )

    assert parsed.parts[0].source_locator == {
        "type": "paragraph",
        "paragraph": 1,
    }

    assert parsed.parts[1].source_locator == {
        "type": "paragraph",
        "paragraph": 3,
    }


def test_parse_docx_document_rejects_empty_content() -> None:
    with pytest.raises(
        ValueError,
        match="Document content must not be empty.",
    ):
        parse_docx_document(b"")


def test_parse_docx_document_rejects_invalid_docx() -> None:
    with pytest.raises(
        ValueError,
        match="DOCX document could not be parsed.",
    ):
        parse_docx_document(b"not-a-docx")


def test_parse_docx_document_rejects_document_without_text() -> None:
    buffer = BytesIO()

    document = DocxDocument()
    document.save(buffer)

    with pytest.raises(
        ValueError,
        match="DOCX document does not contain readable text.",
    ):
        parse_docx_document(
            buffer.getvalue()
        )


def test_chunk_parsed_document_preserves_part_metadata() -> None:
    document = ParsedDocument(
        document_type=SupportedDocumentType.PDF,
        mime_type="application/pdf",
        parts=[
            ParsedDocumentPart(
                text="Treść pierwszej strony.",
                page_number=1,
                section_title="Wprowadzenie",
                source_locator={
                    "type": "page",
                    "page": 1,
                },
            ),
            ParsedDocumentPart(
                text="Treść drugiej strony.",
                page_number=2,
                section_title="Procedura",
                source_locator={
                    "type": "page",
                    "page": 2,
                },
            ),
        ],
    )

    chunks = chunk_parsed_document(document)

    assert len(chunks) == 2

    assert chunks[0].chunk_index == 0
    assert chunks[0].page_number == 1
    assert chunks[0].section_title == "Wprowadzenie"
    assert chunks[0].source_locator == {
        "type": "page",
        "page": 1,
        "part_index": 0,
    }

    assert chunks[1].chunk_index == 1
    assert chunks[1].page_number == 2
    assert chunks[1].section_title == "Procedura"
    assert chunks[1].source_locator == {
        "type": "page",
        "page": 2,
        "part_index": 1,
    }


def test_chunk_parsed_document_reindexes_chunks_across_parts() -> None:
    document = ParsedDocument(
        document_type=SupportedDocumentType.PDF,
        mime_type="application/pdf",
        parts=[
            ParsedDocumentPart(
                text="A" * 30,
                page_number=1,
            ),
            ParsedDocumentPart(
                text="B" * 30,
                page_number=2,
            ),
        ],
    )

    chunks = chunk_parsed_document(
        document,
        config=ChunkingConfig(
            max_characters=20,
            overlap_characters=5,
        ),
    )

    assert [chunk.chunk_index for chunk in chunks] == [
        0,
        1,
        2,
        3,
    ]

    assert [chunk.page_number for chunk in chunks] == [
        1,
        1,
        2,
        2,
    ]