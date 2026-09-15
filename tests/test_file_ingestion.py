import uuid
from io import BytesIO

from app.db.session import SessionLocal
from app.ingestion.file_ingestion import (
    FileIngestionRequest,
    ingest_document_file,
)
from app.models.document_chunk import DocumentChunk
from app.models.document_version import DocumentVersion
from app.models.organization import Organization
from docx import Document as DocxDocument
from sqlalchemy import select


def test_ingest_docx_preserves_section_metadata() -> None:
    organization_id = uuid.uuid4()

    buffer = BytesIO()
    document = DocxDocument()
    document.add_heading(
        "Procedura magazynowa",
        level=1,
    )
    document.add_paragraph(
        "Towar należy sprawdzić przed przyjęciem."
    )
    document.add_heading(
        "Kontrola jakości",
        level=2,
    )
    document.add_paragraph(
        "Należy zweryfikować stan opakowania."
    )
    document.save(buffer)

    with SessionLocal() as db:
        db.add(
            Organization(
                id=organization_id,
                code=f"FILE-{organization_id.hex[:8]}",
                name="File Ingestion Test",
            )
        )
        db.flush()

        result = ingest_document_file(
            db,
            FileIngestionRequest(
                organization_id=organization_id,
                title="Procedura magazynowa",
                source_type="upload",
                filename="procedura.docx",
                content=buffer.getvalue(),
                source_system="upload",
                external_id=f"docx-{uuid.uuid4()}",
            ),
        )

        assert result.created_document is True
        assert result.created_version is True
        assert result.document_version_id is not None
        assert result.chunks_created == 2

        chunks = db.scalars(
            select(DocumentChunk)
            .where(
                DocumentChunk.document_version_id
                == result.document_version_id
            )
            .order_by(DocumentChunk.chunk_index)
        ).all()

        assert len(chunks) == 2

        assert chunks[0].section_title == "Procedura magazynowa"
        assert chunks[0].source_locator == {
            "type": "paragraph",
            "paragraph": 1,
            "part_index": 0,
        }

        assert chunks[1].section_title == "Kontrola jakości"
        assert chunks[1].source_locator == {
            "type": "paragraph",
            "paragraph": 3,
            "part_index": 1,
        }

        db.rollback()


def test_file_ingestion_persists_resolved_mime_type() -> None:
    organization_id = uuid.uuid4()

    with SessionLocal() as db:
        db.add(
            Organization(
                id=organization_id,
                code=f"FILE-MIME-{organization_id.hex[:8]}",
                name="File MIME Test",
            )
        )
        db.flush()

        result = ingest_document_file(
            db,
            FileIngestionRequest(
                organization_id=organization_id,
                title="Instrukcja",
                source_type="upload",
                filename="instrukcja.txt",
                content="Treść instrukcji.".encode(),
                source_system="upload",
                external_id=f"txt-{uuid.uuid4()}",
            ),
        )

        assert result.document_version_id is not None

        version = db.scalar(
            select(DocumentVersion).where(
                DocumentVersion.id
                == result.document_version_id
            )
        )

        assert version is not None
        assert version.mime_type == "text/plain"

        db.rollback()


def test_file_ingestion_does_not_commit_transaction() -> None:
    organization_id = uuid.uuid4()

    with SessionLocal() as db:
        db.add(
            Organization(
                id=organization_id,
                code=f"FILE-TX-{organization_id.hex[:8]}",
                name="File Transaction Test",
            )
        )
        db.flush()

        result = ingest_document_file(
            db,
            FileIngestionRequest(
                organization_id=organization_id,
                title="Instrukcja transakcyjna",
                source_type="upload",
                filename="instrukcja.txt",
                content="Treść dokumentu.".encode(),
                source_system="upload",
                external_id=f"tx-{uuid.uuid4()}",
            ),
        )

        document_version_id = result.document_version_id

        assert document_version_id is not None

        db.rollback()

        stored_version = db.scalar(
            select(DocumentVersion).where(
                DocumentVersion.id == document_version_id
            )
        )

        assert stored_version is None