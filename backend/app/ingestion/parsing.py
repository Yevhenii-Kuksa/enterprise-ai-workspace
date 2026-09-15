from dataclasses import dataclass
from enum import StrEnum
from io import BytesIO

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.ingestion.chunking import (
    ChunkingConfig,
    TextChunk,
    chunk_text,
)


class SupportedDocumentType(StrEnum):
    TXT = "txt"
    PDF = "pdf"
    DOCX = "docx"


@dataclass(frozen=True, slots=True)
class ParsedDocumentPart:
    text: str
    page_number: int | None = None
    section_title: str | None = None
    source_locator: dict[str, object] | None = None


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    document_type: SupportedDocumentType
    mime_type: str
    parts: list[ParsedDocumentPart]

    @property
    def text(self) -> str:
        return "\n\n".join(
            part.text
            for part in self.parts
            if part.text
        )


def parse_txt_document(
    content: bytes,
    *,
    encoding: str = "utf-8",
) -> ParsedDocument:
    if not content:
        raise ValueError("Document content must not be empty.")

    try:
        text = content.decode(encoding)
    except UnicodeDecodeError as exc:
        raise ValueError(
            f"TXT document is not valid {encoding} text."
        ) from exc

    normalized_text = (
        text.replace("\r\n", "\n")
        .replace("\r", "\n")
        .strip()
    )

    if not normalized_text:
        raise ValueError(
            "TXT document does not contain readable text."
        )

    return ParsedDocument(
        document_type=SupportedDocumentType.TXT,
        mime_type="text/plain",
        parts=[
            ParsedDocumentPart(
                text=normalized_text,
                source_locator={
                    "type": "document",
                },
            )
        ],
    )


def parse_pdf_document(
    content: bytes,
) -> ParsedDocument:
    if not content:
        raise ValueError("Document content must not be empty.")

    try:
        reader = PdfReader(BytesIO(content))
    except Exception as exc:
        raise ValueError(
            "PDF document could not be parsed."
        ) from exc

    parts: list[ParsedDocumentPart] = []

    for page_index, page in enumerate(reader.pages):
        page_number = page_index + 1

        try:
            text = page.extract_text() or ""
        except Exception as exc:
            raise ValueError(
                f"Could not extract text from PDF page {page_number}."
            ) from exc

        normalized_text = text.strip()

        if not normalized_text:
            continue

        parts.append(
            ParsedDocumentPart(
                text=normalized_text,
                page_number=page_number,
                source_locator={
                    "type": "page",
                    "page": page_number,
                },
            )
        )

    if not parts:
        raise ValueError(
            "PDF document does not contain extractable text."
        )

    return ParsedDocument(
        document_type=SupportedDocumentType.PDF,
        mime_type="application/pdf",
        parts=parts,
    )


def parse_docx_document(
    content: bytes,
) -> ParsedDocument:
    if not content:
        raise ValueError("Document content must not be empty.")

    try:
        document = DocxDocument(BytesIO(content))
    except Exception as exc:
        raise ValueError(
            "DOCX document could not be parsed."
        ) from exc

    parts: list[ParsedDocumentPart] = []
    current_section_title: str | None = None

    for paragraph_index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()

        if not text:
            continue

        style_name = (
            paragraph.style.name
            if paragraph.style is not None
            else ""
        )

        if style_name.startswith("Heading"):
            current_section_title = text
            continue

        parts.append(
            ParsedDocumentPart(
                text=text,
                section_title=current_section_title,
                source_locator={
                    "type": "paragraph",
                    "paragraph": paragraph_index,
                },
            )
        )

    if not parts:
        raise ValueError(
            "DOCX document does not contain readable text."
        )

    return ParsedDocument(
        document_type=SupportedDocumentType.DOCX,
        mime_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        parts=parts,
    )


def chunk_parsed_document(
    document: ParsedDocument,
    *,
    config: ChunkingConfig | None = None,
) -> list[TextChunk]:
    chunks: list[TextChunk] = []
    next_chunk_index = 0

    for part_index, part in enumerate(document.parts):
        part_chunks = chunk_text(
            part.text,
            config=config,
        )

        for part_chunk in part_chunks:
            source_locator = dict(
                part.source_locator or {}
            )
            source_locator.setdefault(
                "part_index",
                part_index,
            )

            chunks.append(
                TextChunk(
                    chunk_index=next_chunk_index,
                    content=part_chunk.content,
                    content_sha256=part_chunk.content_sha256,
                    token_count=part_chunk.token_count,
                    page_number=part.page_number,
                    section_title=part.section_title,
                    source_locator=source_locator,
                )
            )

            next_chunk_index += 1

    return chunks